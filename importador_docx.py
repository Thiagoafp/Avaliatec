"""
importador_docx.py — Lê provas e gabaritos em .docx

FORMATOS SUPORTADOS:
────────────────────────────────────────────────────────────────────

1) QUESTÃO V/F  (formato SENAI com afirmativas numeradas)
   Prova:
       1. Enunciado da questão
       I.  (     )  Texto da afirmativa I
       II. (     )  Texto da afirmativa II

   Gabarito separado:
       Questão 1: ...
       I: [F]  comentário
       II: [V]  comentário

2) QUESTÃO DISSERTATIVA embutida na prova VF
   (detectada quando a questão começa com número mas NÃO tem afirmativas)
   Gabarito da dissertativa no arquivo de gabarito:
       Questão 7 (30.0 pts):
       Resposta esperada: texto...

3) QUESTÃO OBJETIVA (múltipla escolha)
   QUESTÃO [objetiva] [2.0]
   Enunciado
   *a) Correta
   b) Errada

4) QUESTÃO DISSERTATIVA standalone
   QUESTÃO [dissertativa] [3.0]
   Enunciado
   GABARITO: texto
"""

import re
from docx import Document


# ─── UTILITÁRIOS ──────────────────────────────────────────────────────────────

def _parse_peso_instrucoes(linhas: list) -> float:
    """Extrai peso unitário das instruções, ex: '6×6.7pts'."""
    for linha in linhas[:10]:
        m = re.search(r'(\d+)\s*[×x\*]\s*([0-9.,]+)\s*pts?', linha, re.IGNORECASE)
        if m:
            try:
                return float(m.group(2).replace(',', '.'))
            except Exception:
                pass
    return 1.0


def _parse_peso_dissertativa_instrucoes(linhas: list) -> float:
    """Extrai peso das dissertativas das instruções, ex: 'Dissertativas (1×30.0pts)'."""
    for linha in linhas[:10]:
        m = re.search(r'Dissertativas?\s*\((\d+)\s*[×x]\s*([0-9.,]+)', linha, re.IGNORECASE)
        if m:
            try:
                return float(m.group(2).replace(',', '.'))
            except Exception:
                pass
    return 10.0


_ROMANOS = re.compile(
    r'^(I{1,3}|IV|VI{0,3}|IX|XI{0,3}|XIV|XV|XVI{0,3})\.\s+\(\s*\)\s+(.+)$'
)

def _is_afirmativa(linha: str):
    m = _ROMANOS.match(linha)
    return (m.group(1), m.group(2)) if m else None


def _is_cabecalho_secao(linha: str) -> bool:
    """Detecta cabeçalhos como 'Questões Dissertativas (1 questões × 30.0 pts cada)'."""
    low = linha.lower()
    return bool(re.search(r'questões?\s+(dissertativas?|objetivas?|vf|verdadeiro)', low))


# ─── PARSER DE PROVA SENAI (VF + Dissertativa misto) ────────────────────────

def _parse_numero_questao(linha: str):
    """Detecta '7. Texto' mas não confunde com afirmativas romanas."""
    m = re.match(r'^(\d+)\.\s+(.+)$', linha)
    return (int(m.group(1)), m.group(2)) if m else None


def importar_prova_senai(caminho: str) -> list:
    """
    Importa prova no formato SENAI (VF + dissertativas misturadas).
    Detecta tipo por presença/ausência de afirmativas.
    """
    doc = Document(caminho)
    linhas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    peso_vf   = _parse_peso_instrucoes(linhas)
    peso_dis  = _parse_peso_dissertativa_instrucoes(linhas)

    questoes  = []
    q_atual   = None

    for linha in linhas:
        # Ignora cabeçalhos de seção
        if _is_cabecalho_secao(linha):
            continue

        nq = _parse_numero_questao(linha)
        if nq:
            if q_atual is not None:
                questoes.append(q_atual)
            q_atual = {
                'numero':   nq[0],
                'enunciado': nq[1],
                'tipo':     None,       # definido ao final
                'peso':     None,       # definido ao final
                'afirmativas': [],
                'gabarito_dissertativa': None,
            }
            continue

        af = _is_afirmativa(linha)
        if af and q_atual is not None:
            q_atual['afirmativas'].append({
                'letra':   af[0],
                'texto':   af[1],
                'gabarito': None,
            })
            continue

    if q_atual is not None:
        questoes.append(q_atual)

    # Define tipo e peso conforme presença de afirmativas
    for q in questoes:
        if q['afirmativas']:
            q['tipo']  = 'vf'
            q['peso']  = peso_vf
        else:
            q['tipo']  = 'dissertativa'
            q['peso']  = peso_dis

    return questoes


# ─── PARSER DE GABARITO SENAI ─────────────────────────────────────────────────

def importar_gabarito_senai(caminho: str) -> dict:
    """
    Importa gabarito SENAI.
    Retorna:
    {
      num_questao: {
        'tipo': 'vf' | 'dissertativa',
        'afirmativas': {letra: 'V'/'F'},   # para vf
        'gabarito_texto': str,             # para dissertativa
        'peso': float,
      }
    }
    """
    doc = Document(caminho)
    linhas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    gabarito = {}
    q_num    = None
    q_tipo   = None
    lendo_resposta = False

    for linha in linhas:
        # "Questão 7 (30.0 pts):" ou "Questão 1: texto"
        m = re.match(r'^Questão\s+(\d+)(?:\s*\(([0-9.,]+)\s*pts?\))?', linha, re.IGNORECASE)
        if m:
            q_num  = int(m.group(1))
            peso_g = float(m.group(2).replace(',', '.')) if m.group(2) else None
            gabarito[q_num] = {
                'tipo': None,
                'afirmativas': {},
                'gabarito_texto': '',
                'peso': peso_g,
            }
            lendo_resposta = False
            # Determina tipo pela presença de "(pts)" → dissertativa
            q_tipo = 'dissertativa' if m.group(2) else 'vf'
            gabarito[q_num]['tipo'] = q_tipo
            continue

        if q_num is None:
            continue

        # Afirmativa VF: "I: [F]  comentário"
        m = re.match(r'^(I{1,3}|IV|VI{0,3}|IX|V)\s*:\s*\[([VF])\]', linha)
        if m:
            gabarito[q_num]['afirmativas'][m.group(1)] = m.group(2)
            gabarito[q_num]['tipo'] = 'vf'
            lendo_resposta = False
            continue

        # Início da resposta dissertativa
        if re.match(r'^Resposta\s+esperada\s*:', linha, re.IGNORECASE):
            texto = re.sub(r'^Resposta\s+esperada\s*:\s*', '', linha, flags=re.IGNORECASE)
            gabarito[q_num]['gabarito_texto'] = texto
            gabarito[q_num]['tipo'] = 'dissertativa'
            lendo_resposta = True
            continue

        # Continua texto da resposta dissertativa (linhas seguintes)
        if lendo_resposta and q_tipo == 'dissertativa':
            # Para se encontrar nova questão ou seção
            if re.match(r'^(Dica|Critério|Questão|\*)', linha, re.IGNORECASE):
                lendo_resposta = False
            else:
                gabarito[q_num]['gabarito_texto'] += ' ' + linha

    return gabarito


def aplicar_gabarito(questoes: list, gabarito: dict) -> list:
    """Aplica o gabarito nas questões."""
    for q in questoes:
        g = gabarito.get(q['numero'])
        if not g:
            continue
        if q['tipo'] == 'vf':
            for af in q['afirmativas']:
                af['gabarito'] = g['afirmativas'].get(af['letra'])
        elif q['tipo'] == 'dissertativa':
            q['gabarito_dissertativa'] = g.get('gabarito_texto', '').strip()
    return questoes


# ─── PARSER OBJETIVA/DISSERTATIVA STANDALONE ─────────────────────────────────

def _parse_peso_header(linha: str) -> float:
    matches = re.findall(r'\[([0-9]+[.,]?[0-9]*)\]', linha)
    for m in matches:
        try:
            return float(m.replace(',', '.'))
        except ValueError:
            continue
    return 1.0


def importar_prova_objetiva(caminho: str) -> list:
    """Importa prova no formato QUESTÃO [objetiva/dissertativa] [peso]."""
    doc = Document(caminho)
    linhas = [p.text.strip() for p in doc.paragraphs]

    questoes = []
    q_atual  = None
    env_lines = []

    for linha in linhas:
        if not linha:
            continue

        if re.match(r'^QUEST[ÃA]O', linha, re.IGNORECASE):
            if q_atual is not None:
                q_atual['enunciado'] = ' '.join(env_lines).strip()
                questoes.append(q_atual)
            tipo = 'dissertativa' if 'dissertativa' in linha.lower() else 'objetiva'
            q_atual = {
                'tipo': tipo,
                'peso': _parse_peso_header(linha),
                'enunciado': '',
                'gabarito_dissertativa': None,
                'alternativas': [],
                'afirmativas': [],
            }
            env_lines = []
            continue

        if q_atual is None:
            continue

        m = re.match(r'^\*?([a-eA-E])[)\.]\s+(.+)$', linha.strip())
        if m:
            correta = linha.strip().startswith('*')
            q_atual['alternativas'].append({
                'letra': m.group(1).lower(),
                'texto': m.group(2).strip(),
                'correta': correta,
            })
            continue

        if linha.upper().startswith('GABARITO:'):
            q_atual['gabarito_dissertativa'] = linha[9:].strip()
            continue

        env_lines.append(linha)

    if q_atual is not None:
        q_atual['enunciado'] = ' '.join(env_lines).strip()
        questoes.append(q_atual)

    for i, q in enumerate(questoes, 1):
        q['numero'] = i

    return questoes


# ─── DETECTOR DE FORMATO ─────────────────────────────────────────────────────

def detectar_formato(caminho: str) -> str:
    doc = Document(caminho)
    linhas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for linha in linhas:
        if re.match(r'^QUEST[ÃA]O', linha, re.IGNORECASE):
            return 'objetiva'
        if re.match(r'^\d+\.\s+', linha):
            return 'senai'
    return 'senai'


# ─── INTERFACE UNIFICADA ─────────────────────────────────────────────────────

def importar_docx(caminho: str, caminho_gabarito: str = None) -> list:
    fmt = detectar_formato(caminho)

    if fmt == 'senai':
        questoes = importar_prova_senai(caminho)
        if caminho_gabarito:
            gabarito = importar_gabarito_senai(caminho_gabarito)
            questoes = aplicar_gabarito(questoes, gabarito)
    else:
        questoes = importar_prova_objetiva(caminho)

    return questoes


def validar_questoes(questoes: list, **kwargs) -> list:
    erros = []
    for i, q in enumerate(questoes, start=1):
        if not q.get('enunciado'):
            erros.append(f"Questão {i}: enunciado vazio.")

        if q.get('tipo') == 'vf':
            if not q.get('afirmativas'):
                erros.append(f"Questão {i} (V/F): nenhuma afirmativa encontrada.")
            sem_gab = [af['letra'] for af in q.get('afirmativas', [])
                       if af.get('gabarito') not in ('V', 'F')]
            if sem_gab:
                erros.append(
                    f"Questão {i} (V/F): afirmativas sem gabarito: {', '.join(sem_gab)}. "
                    "Importe o arquivo de gabarito."
                )

        elif q.get('tipo') == 'objetiva':
            if len(q.get('alternativas', [])) < 2:
                erros.append(f"Questão {i} (objetiva): menos de 2 alternativas.")
            corretas = [a for a in q.get('alternativas', []) if a['correta']]
            if len(corretas) != 1:
                erros.append(
                    f"Questão {i} (objetiva): marque exatamente 1 alternativa correta com *. "
                    f"Encontradas: {len(corretas)}."
                )

    return erros
