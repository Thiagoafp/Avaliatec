"""
exportar.py — Geração de Excel de resultados de provas
Abas: Resumo da Turma | Respostas Detalhadas | Análise por Questão | Para Impressão
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── ESTILOS ──────────────────────────────────────────────────────────────────
COR_SENAI     = "C0000F"
COR_HDR_FG    = "FFFFFF"
COR_NAO_LOGOU = "FFCCCC"
COR_NAO_ENVIOU= "FFF3CD"
COR_ENVIADA   = "D4EDDA"
COR_TEMPO     = "FFE0B2"
COR_ACERTO    = "C8E6C9"
COR_ERRO      = "FFCDD2"
COR_ALT0      = "FFFFFF"
COR_ALT1      = "F5F5F5"
COR_TOTAL     = "EEEEEE"

def _fill(c): return PatternFill("solid", fgColor=c)
def _font(bold=False, size=10, color="000000", italic=False):
    return Font(bold=bold, size=size, color=color, italic=italic, name="Arial")

_borda = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC'),
)
_borda_media = Border(
    left=Side(style='medium', color='AAAAAA'),
    right=Side(style='medium', color='AAAAAA'),
    top=Side(style='medium', color='AAAAAA'),
    bottom=Side(style='medium', color='AAAAAA'),
)

def _hdr(ws, row, col, value, bg="555555", fg="FFFFFF", bold=True,
         size=10, align="center", wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=size, color=fg, name="Arial")
    c.fill = _fill(bg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = _borda
    return c

def _cel(ws, row, col, value, bg="FFFFFF", bold=False, size=10,
         align="left", wrap=False, color="000000"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=size, color=color, name="Arial")
    c.fill = _fill(bg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = _borda
    return c


# ─── GERAR EXCEL COMPLETO ─────────────────────────────────────────────────────

def gerar_excel(prova: dict, dados: list, questoes_analise: list = None) -> bytes:
    """
    prova             — dict da prova
    dados             — lista de db.gerar_dados_export()
    questoes_analise  — lista de dicts com análise por questão (opcional)
    """
    wb = Workbook()

    _aba_resumo(wb, prova, dados)
    _aba_respostas(wb, dados)
    if questoes_analise:
        _aba_analise(wb, prova, questoes_analise)
    _aba_impressao(wb, dados)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─── ABA 1: RESUMO DA TURMA ───────────────────────────────────────────────────

def _aba_resumo(wb, prova, dados):
    ws = wb.active
    ws.title = "Resumo da Turma"

    # Título
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"SENAI FATESG — {prova['titulo']}"
    c.font = Font(bold=True, size=14, color="FFFFFF", name="Arial")
    c.fill = _fill(COR_SENAI)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Subtítulo
    ws.merge_cells("A2:H2")
    info = (f"Tempo: {prova['tempo_minutos']} min  |  "
            f"Status: {prova['status']}  |  "
            f"Gerado em: {__import__('datetime').datetime.now().strftime('%d/%m/%Y %H:%M')}")
    c = ws["A2"]
    c.value = info
    c.font = Font(italic=True, size=9, color="555555", name="Arial")
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 16

    # Cabeçalhos
    hdrs = ["Nº", "Nome do Aluno", "Status", "Nota /10",
            "Acertos", "Erros", "Respondidas", "Enviada em"]
    for ci, h in enumerate(hdrs, 1):
        _hdr(ws, 3, ci, h)
    ws.row_dimensions[3].height = 20

    # Larguras
    for col, w in zip("ABCDEFGH", [5, 38, 30, 10, 10, 10, 14, 20]):
        ws.column_dimensions[col].width = w

    # Dados
    alunos_reais = [d for d in dados if not d['Nome'].endswith(' *')]
    alunos_extra = [d for d in dados if d['Nome'].endswith(' *')]
    notas_validas = []

    for i, linha in enumerate(alunos_reais + alunos_extra, start=4):
        alt = COR_ALT1 if i % 2 == 0 else COR_ALT0
        status_str = str(linha.get('Status', ''))

        # Cor de status
        s = status_str.lower()
        if 'não logou' in s:        sf = COR_NAO_LOGOU
        elif 'não enviou' in s:     sf = COR_NAO_ENVIOU
        elif 'tempo esgotado' in s: sf = COR_TEMPO
        elif 'enviada' in s or 'nota' in s: sf = COR_ENVIADA
        else:                       sf = alt

        nota = linha.get('Nota', '')
        try:
            nota_f = float(nota) if nota not in ('', None) else None
        except Exception:
            nota_f = None

        if nota_f is not None:
            notas_validas.append(nota_f)
            nf = COR_ACERTO if nota_f >= 5 else COR_ERRO
        else:
            nf = alt

        # Acertos e erros a partir das colunas de questão
        q_cols = [k for k in linha if k.startswith('Q') and 'pt' in k]
        acertos = sum(1 for k in q_cols if '✓' in str(linha.get(k, '')))
        erros   = sum(1 for k in q_cols if '✗' in str(linha.get(k, '')))

        _cel(ws, i, 1, i - 3,     bg=alt,  align="center")
        _cel(ws, i, 2, linha['Nome'], bg=alt, bold=True)
        _cel(ws, i, 3, status_str, bg=sf)
        _cel(ws, i, 4, nota_f if nota_f is not None else '', bg=nf,
             bold=True, align="center",
             color="198754" if nota_f and nota_f >= 5 else
                   "C0000F" if nota_f and nota_f < 5 else "000000")
        _cel(ws, i, 5, acertos if q_cols else '',   bg=COR_ACERTO if acertos else alt, align="center")
        _cel(ws, i, 6, erros   if q_cols else '',   bg=COR_ERRO   if erros   else alt, align="center")
        _cel(ws, i, 7, linha.get('Respondidas', ''), bg=alt, align="center")
        _cel(ws, i, 8, linha.get('Enviada em', ''),  bg=alt, size=9)
        ws.row_dimensions[i].height = 16

    # Linha de totais
    r = len(dados) + 4
    ws.merge_cells(f"A{r}:B{r}")
    c = ws[f"A{r}"]
    c.value = f"Total: {len(dados)} aluno(s)"
    c.font = Font(bold=True, name="Arial")
    c.fill = _fill(COR_TOTAL)

    nao_logou  = sum(1 for d in dados if 'Não logou' in str(d.get('Status','')))
    nao_enviou = sum(1 for d in dados if 'não enviou' in str(d.get('Status','')).lower())
    enviadas   = len(dados) - nao_logou - nao_enviou

    ws[f"C{r}"] = f"Enviadas: {enviadas} | Não enviou: {nao_enviou} | Ausentes: {nao_logou}"
    ws[f"C{r}"].fill = _fill(COR_TOTAL)
    ws[f"C{r}"].font = Font(italic=True, size=9, name="Arial")

    if notas_validas:
        media = sum(notas_validas) / len(notas_validas)
        maior = max(notas_validas)
        menor = min(notas_validas)
        ws[f"D{r}"] = f"Média: {media:.1f}"
        ws[f"D{r}"].font = Font(bold=True, name="Arial")
        ws[f"D{r}"].fill = _fill(COR_TOTAL)
        ws[f"E{r}"] = f"Maior: {maior:.1f}"
        ws[f"E{r}"].fill = _fill(COR_TOTAL)
        ws[f"F{r}"] = f"Menor: {menor:.1f}"
        ws[f"F{r}"].fill = _fill(COR_TOTAL)


# ─── ABA 2: RESPOSTAS DETALHADAS ──────────────────────────────────────────────

def _aba_respostas(wb, dados):
    ws = wb.create_sheet("Respostas Detalhadas")

    base = ["Nome", "Status", "Nota"]
    q_cols = [k for k in (dados[0].keys() if dados else [])
              if k.startswith('Q') and 'pt' in k]
    all_cols = base + q_cols

    for ci, h in enumerate(all_cols, 1):
        _hdr(ws, 1, ci, h, wrap=True)
    ws.row_dimensions[1].height = 32
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    for ci in range(4, len(all_cols) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 20

    for i, linha in enumerate(dados, start=2):
        alt = COR_ALT1 if i % 2 == 0 else COR_ALT0
        for ci, col in enumerate(all_cols, 1):
            val = linha.get(col, '')
            v = str(val) if val not in ('', None) else ''
            if ci <= 3:
                bg = alt
            elif '✓' in v: bg = COR_ACERTO
            elif '✗' in v: bg = COR_ERRO
            else:          bg = alt
            _cel(ws, i, ci, v, bg=bg, wrap=(ci > 3), size=9)
        ws.row_dimensions[i].height = 28


# ─── ABA 3: ANÁLISE POR QUESTÃO ───────────────────────────────────────────────

def _aba_analise(wb, prova, questoes_analise):
    """
    questoes_analise = lista de dicts:
    {
      'numero': int,
      'enunciado': str,
      'tipo': str,
      'peso': float,
      'total_alunos': int,    # que fizeram a prova
      'acertos': int/float,
      'erros': int/float,
      'taxa_acerto': float,   # 0-100
      # Para VF: lista de afirmativas
      'afirmativas': [{'letra','texto','gabarito','acertos','erros','taxa'}]
    }
    """
    ws = wb.create_sheet("Análise por Questão")

    # Título
    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = f"Análise por Questão — {prova['titulo']}"
    c.font = Font(bold=True, size=13, color="FFFFFF", name="Arial")
    c.fill = _fill(COR_SENAI)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Cabeçalhos
    hdrs = ["Questão", "Tipo", "Peso", "Enunciado", "Acertos", "Erros", "% Acerto"]
    for ci, h in enumerate(hdrs, 1):
        _hdr(ws, 2, ci, h)
    ws.row_dimensions[2].height = 18

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 55
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 12

    r = 3
    for q in questoes_analise:
        alt = COR_ALT1 if r % 2 == 0 else COR_ALT0
        taxa = q.get('taxa_acerto', 0)
        cor_taxa = COR_ACERTO if taxa >= 60 else COR_TEMPO if taxa >= 40 else COR_ERRO

        _cel(ws, r, 1, f"Q{q['numero']}", bg=alt, bold=True, align="center")
        _cel(ws, r, 2, q['tipo'].upper(), bg=alt, align="center", size=9)
        _cel(ws, r, 3, q['peso'],         bg=alt, align="center")
        _cel(ws, r, 4, q['enunciado'][:120], bg=alt, wrap=True, size=9)
        _cel(ws, r, 5, q.get('acertos', ''), bg=COR_ACERTO if q.get('acertos') else alt,
             align="center", bold=True)
        _cel(ws, r, 6, q.get('erros', ''),   bg=COR_ERRO   if q.get('erros')   else alt,
             align="center", bold=True)
        _cel(ws, r, 7, f"{taxa:.1f}%",       bg=cor_taxa,
             bold=True, align="center",
             color="155724" if taxa >= 60 else "856404" if taxa >= 40 else "842029")
        ws.row_dimensions[r].height = 18
        r += 1

        # Afirmativas VF — sublinhas indentadas
        for af in q.get('afirmativas', []):
            taxa_af = af.get('taxa', 0)
            cor_af  = COR_ACERTO if taxa_af >= 60 else COR_TEMPO if taxa_af >= 40 else COR_ERRO
            gab     = af.get('gabarito', '?')
            _cel(ws, r, 1, f"  {af['letra']}",   bg="F8F8F8", align="center", size=9)
            _cel(ws, r, 2, f"[{gab}]",            bg="F8F8F8", align="center", size=9,
                 color="198754" if gab == 'V' else "C0000F")
            _cel(ws, r, 3, '',                    bg="F8F8F8")
            _cel(ws, r, 4, f"   {af['texto'][:110]}", bg="F8F8F8", wrap=True, size=9)
            _cel(ws, r, 5, af.get('acertos', ''), bg=COR_ACERTO if af.get('acertos') else "F8F8F8",
                 align="center", size=9)
            _cel(ws, r, 6, af.get('erros', ''),   bg=COR_ERRO   if af.get('erros')   else "F8F8F8",
                 align="center", size=9)
            _cel(ws, r, 7, f"{taxa_af:.1f}%",     bg=cor_af, align="center", size=9)
            ws.row_dimensions[r].height = 15
            r += 1

    # Linha de média geral
    ws.merge_cells(f"A{r}:C{r}")
    c = ws[f"A{r}"]
    c.value = "Média geral de acertos"
    c.font = Font(bold=True, name="Arial")
    c.fill = _fill(COR_TOTAL)

    taxas = [q.get('taxa_acerto', 0) for q in questoes_analise if q.get('taxa_acerto') is not None]
    if taxas:
        media_taxa = sum(taxas) / len(taxas)
        cor_m = COR_ACERTO if media_taxa >= 60 else COR_TEMPO if media_taxa >= 40 else COR_ERRO
        _cel(ws, r, 7, f"{media_taxa:.1f}%", bg=cor_m, bold=True, align="center")


# ─── ABA 4: PARA IMPRESSÃO ────────────────────────────────────────────────────

def _aba_impressao(wb, dados):
    ws = wb.create_sheet("Para Impressão")
    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = "FICHA DE RESPOSTAS — Impressão individual por aluno"
    c.font = Font(bold=True, italic=True, color="555555", name="Arial")
    c.alignment = Alignment(horizontal="center")

    q_cols = [k for k in (dados[0].keys() if dados else [])
              if k.startswith('Q') and 'pt' in k]

    r = 3
    for linha in dados:
        if linha.get('Status', '') in ('Não logou', ''):
            continue
        nota = linha.get('Nota', '—')
        ws.merge_cells(f"A{r}:C{r}")
        c = ws[f"A{r}"]
        c.value = (f"Aluno: {linha['Nome']}   |   "
                   f"Nota: {nota}   |   "
                   f"Status: {linha.get('Status','')}")
        c.font = Font(bold=True, name="Arial")
        c.fill = _fill("EEEEEE")
        r += 1
        for col in q_cols:
            _cel(ws, r, 1, col,                    bg="F5F5F5", bold=True, size=9)
            _cel(ws, r, 2, str(linha.get(col,'—')), bg="FFFFFF", wrap=True, size=9)
            r += 1
        r += 1

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
